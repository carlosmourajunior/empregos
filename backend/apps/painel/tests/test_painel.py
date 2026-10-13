from io import BytesIO

import pytest
from django.contrib.auth.models import Permission
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.contas.models import Usuario
from apps.curriculos.models import Curriculo
from apps.vagas.models import Interesse, Vaga


def criar_vaga(empresa, **extra):
    dados = dict(
        cargo="Auxiliar de padeiro",
        area="Cozinha",
        descricao="Ajudar na padaria.",
        contrato="clt",
        horario="manha",
        bairro="Centro",
        quantidade=2,
    )
    return Vaga.objects.create(empresa=empresa, **{**dados, **extra})


@pytest.fixture
def aprovador(db):
    usuario = Usuario.objects.create_user("39053344705", "Ana da Prefeitura", "5511966665555", "segredo123")
    usuario.tipo = Usuario.Tipo.PREFEITURA
    usuario.save()
    usuario.user_permissions.add(Permission.objects.get(codename="aprovar_vaga"))
    return usuario


@pytest.fixture
def painel(api, aprovador):
    api.force_authenticate(aprovador)
    return api


def test_so_aprovador_entra(api, candidato, empresa):
    for usuario in (candidato, empresa):
        api.force_authenticate(usuario)
        assert api.get("/api/painel/vagas/").status_code == 403
        assert api.get("/api/painel/numeros/").status_code == 403


def test_fila_de_aprovacao_mais_antiga_primeiro(painel, empresa):
    primeira = criar_vaga(empresa, cargo="Primeira")
    criar_vaga(empresa, cargo="Segunda")
    criar_vaga(empresa, cargo="Aprovada", status=Vaga.Status.APROVADA)
    resposta = painel.get("/api/painel/vagas/?status=pendente").json()
    assert [v["cargo"] for v in resposta["results"]] == ["Primeira", "Segunda"]
    assert resposta["results"][0]["id"] == primeira.id
    assert resposta["results"][0]["empresa_cnpj"] == "11222333000181"


def test_aprovar_varias_registra_quem_aprovou(painel, aprovador, empresa):
    vagas = [criar_vaga(empresa), criar_vaga(empresa)]
    resposta = painel.post("/api/painel/vagas/avaliar/", {"ids": [v.id for v in vagas], "decisao": "aprovar"})
    assert resposta.json() == {"alteradas": 2}
    for vaga in vagas:
        vaga.refresh_from_db()
        assert vaga.status == Vaga.Status.APROVADA
        assert vaga.avaliada_por == aprovador
        assert vaga.avaliada_em is not None


def test_recusar_exige_motivo_e_empresa_ve(painel, api, empresa):
    vaga = criar_vaga(empresa)
    sem_motivo = painel.post("/api/painel/vagas/avaliar/", {"ids": [vaga.id], "decisao": "recusar"}, format="json")
    assert sem_motivo.status_code == 400
    painel.post(
        "/api/painel/vagas/avaliar/",
        {"ids": [vaga.id], "decisao": "recusar", "motivo": "  Salário abaixo do mínimo "},
        format="json",
    )
    api.force_authenticate(empresa)
    minha = api.get("/api/minhas-vagas/").json()[0]
    assert minha["status"] == "recusada"
    assert minha["motivo_recusa"] == "Salário abaixo do mínimo"

    # Ao corrigir, a vaga volta para a fila e o motivo some.
    api.patch(f"/api/minhas-vagas/{vaga.id}/", {"salario": "1600"}, format="json")
    vaga.refresh_from_db()
    assert vaga.status == Vaga.Status.PENDENTE
    assert vaga.motivo_recusa == ""


def test_vaga_encerrada_nao_e_reaberta(painel, empresa):
    vaga = criar_vaga(empresa, status=Vaga.Status.ENCERRADA)
    resposta = painel.post("/api/painel/vagas/avaliar/", {"ids": [vaga.id], "decisao": "aprovar"}, format="json")
    assert resposta.json() == {"alteradas": 0}


def test_bloquear_empresa_tira_vagas_do_ar(painel, empresa):
    publico = APIClient()
    criar_vaga(empresa, status=Vaga.Status.APROVADA)
    assert publico.get("/api/vagas/").json()["count"] == 1
    painel.post(f"/api/painel/empresas/{empresa.id}/bloquear/", {"bloqueada": True}, format="json")
    assert publico.get("/api/vagas/").json()["count"] == 0
    assert publico.post("/api/auth/login/", {"documento": "11222333000181", "senha": "segredo123"}).status_code == 400

    painel.post(f"/api/painel/empresas/{empresa.id}/bloquear/", {"bloqueada": False}, format="json")
    assert publico.get("/api/vagas/").json()["count"] == 1


def test_numeros(painel, empresa, candidato):
    vaga = criar_vaga(empresa, status=Vaga.Status.APROVADA, quantidade=3)
    criar_vaga(empresa)
    Curriculo.objects.create(
        usuario=candidato,
        data_nascimento="1990-01-01",
        bairro="Centro",
        escolaridade="medio",
        areas_interesse=["Cozinha"],
    )
    Interesse.objects.create(vaga=vaga, candidato=candidato)
    dados = painel.get("/api/painel/numeros/").json()
    assert dados["pendentes"] == 1
    assert dados["vagas_abertas"] == 1
    assert dados["postos_abertos"] == 3
    assert dados["curriculos"] == 1
    assert dados["interesses"] == 1
    assert dados["por_area"] == [{"area": "Cozinha", "vagas": 3, "interesses": 1, "curriculos": 1}]


def test_equipe_incluir_e_remover(painel, aprovador, candidato):
    assert [m["nome"] for m in painel.get("/api/painel/equipe/").json()] == ["Ana da Prefeitura"]
    assert painel.post("/api/painel/equipe/", {"documento": "000.000.000-00"}).status_code == 400

    resposta = painel.post("/api/painel/equipe/", {"documento": "529.982.247-25"})
    assert resposta.status_code == 201
    candidato.refresh_from_db()
    assert candidato.tipo == "prefeitura"
    assert Usuario.objects.get(pk=candidato.pk).has_perm("vagas.aprovar_vaga")

    assert painel.delete(f"/api/painel/equipe/{aprovador.id}/").status_code == 400
    assert painel.delete(f"/api/painel/equipe/{candidato.id}/").status_code == 204
    candidato = Usuario.objects.get(pk=candidato.pk)
    assert candidato.tipo == "candidato"
    assert not candidato.has_perm("vagas.aprovar_vaga")


@pytest.mark.parametrize("tipo", ["vagas", "curriculos", "empresas"])
def test_planilhas(painel, empresa, candidato, tipo):
    criar_vaga(empresa)
    Curriculo.objects.create(
        usuario=candidato, data_nascimento="1990-01-01", bairro="Centro", escolaridade="medio", areas_interesse=["a"]
    )
    resposta = painel.get(f"/api/painel/exportar/{tipo}/")
    assert resposta.status_code == 200
    assert f"meuemprego-{tipo}-" in resposta["Content-Disposition"]
    folha = load_workbook(BytesIO(resposta.content)).active
    assert folha.max_row == 2
    if tipo == "curriculos":
        assert "52998224725" not in str(list(folha.values))


def test_editar_vaga_mantem_situacao(painel, empresa):
    vaga = criar_vaga(empresa, status=Vaga.Status.APROVADA)
    resposta = painel.patch(
        f"/api/painel/vagas/{vaga.id}/", {"cargo": "  Padeiro  ", "salario": "1800", "horario": "tarde"}, format="json"
    )
    assert resposta.status_code == 200
    assert resposta.json()["cargo"] == "Padeiro"
    assert resposta.json()["horario_nome"] == "Tarde"
    vaga.refresh_from_db()
    assert vaga.status == Vaga.Status.APROVADA
    assert vaga.salario == 1800


def test_editar_vaga_valida_campos(painel, empresa):
    vaga = criar_vaga(empresa)
    resposta = painel.patch(f"/api/painel/vagas/{vaga.id}/", {"quantidade": 0}, format="json")
    assert resposta.status_code == 400
    assert "quantidade" in resposta.json()


def test_editar_empresa_e_novo_telefone_precisa_confirmar(painel, empresa):
    resposta = painel.patch(
        f"/api/painel/empresas/{empresa.id}/",
        {"nome": "Padaria Pão Quente", "telefone": "(11) 95555-4444"},
        format="json",
    )
    assert resposta.status_code == 200
    assert resposta.json()["nome"] == "Padaria Pão Quente"
    empresa.refresh_from_db()
    assert empresa.telefone == "5511955554444"
    assert not empresa.telefone_verificado


def test_telefone_invalido(painel, empresa):
    resposta = painel.patch(f"/api/painel/empresas/{empresa.id}/", {"telefone": "123"}, format="json")
    assert resposta.status_code == 400
    assert "telefone" in resposta.json()


def test_editar_empresa_nao_altera_candidato(painel, candidato):
    assert painel.patch(f"/api/painel/empresas/{candidato.id}/", {"nome": "X"}, format="json").status_code == 404


def test_lista_e_busca_de_candidatos(painel, candidato):
    Curriculo.objects.create(
        usuario=candidato,
        data_nascimento="1990-05-01",
        bairro="Vila Nova",
        escolaridade="medio",
        areas_interesse=["Cozinha"],
    )
    for busca in ("maria", "vila", "529.982", "98888"):
        resultados = painel.get(f"/api/painel/candidatos/?busca={busca}").json()["results"]
        assert [c["nome"] for c in resultados] == ["Maria"], busca
        assert resultados[0]["curriculo"]["bairro"] == "Vila Nova"
    assert painel.get("/api/painel/candidatos/?busca=centro").json()["results"] == []


def test_editar_candidato_e_curriculo(painel, candidato):
    curriculo = {
        "data_nascimento": "1990-05-01",
        "bairro": "Centro",
        "escolaridade": "tecnico",
        "areas_interesse": ["Limpeza", "Cozinha"],
        "tem_cnh": True,
        "categoria_cnh": "B",
        "experiencias": [{"o_que_fazia": "Cozinheira", "onde": "Restaurante", "quanto_tempo": "2 anos"}],
    }
    resposta = painel.patch(
        f"/api/painel/candidatos/{candidato.id}/", {"nome": "Maria Souza", "curriculo": curriculo}, format="json"
    )
    assert resposta.status_code == 200, resposta.json()
    dados = resposta.json()
    assert dados["nome"] == "Maria Souza"
    assert dados["curriculo"]["escolaridade"] == "tecnico"
    assert dados["curriculo"]["experiencias"][0]["onde"] == "Restaurante"

    curriculo["experiencias"] = []
    painel.patch(f"/api/painel/candidatos/{candidato.id}/", {"curriculo": curriculo}, format="json")
    candidato.refresh_from_db()
    assert candidato.curriculo.experiencias.count() == 0
    assert candidato.telefone_verificado


def test_curriculo_invalido_nao_salva_nada(painel, candidato):
    resposta = painel.patch(
        f"/api/painel/candidatos/{candidato.id}/",
        {"nome": "Outra", "curriculo": {"bairro": "Centro"}},
        format="json",
    )
    assert resposta.status_code == 400
    candidato.refresh_from_db()
    assert candidato.nome == "Maria"
