from collections import Counter
from io import BytesIO

from django.contrib.auth.models import Permission
from django.db.models import Count, Q
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from rest_framework import generics, status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.generics import get_object_or_404
from rest_framework.response import Response

from apps.contas.models import Usuario
from apps.curriculos.models import Curriculo
from apps.vagas.models import Interesse, Vaga

from .permissions import PodeAprovar
from .serializers import (
    AvaliarSerializer,
    EmpresaPainelSerializer,
    MembroSerializer,
    NovoMembroSerializer,
    VagaPainelSerializer,
)


class VagasPainelView(generics.ListAPIView):
    """Todas as vagas para a prefeitura. Filtros: ?status= ?busca= ?empresa="""

    permission_classes = [PodeAprovar]
    serializer_class = VagaPainelSerializer

    def get_queryset(self):
        vagas = Vaga.objects.select_related("empresa", "avaliada_por").annotate(interessados=Count("interesses"))
        params = self.request.query_params
        if situacao := params.get("status"):
            vagas = vagas.filter(status=situacao)
        if empresa := params.get("empresa"):
            vagas = vagas.filter(empresa_id=empresa)
        if busca := params.get("busca", "").strip():
            vagas = vagas.filter(
                Q(cargo__icontains=busca) | Q(area__icontains=busca) | Q(empresa__nome__icontains=busca)
            )
        # Na fila de aprovação, a mais antiga primeiro: quem esperou mais é atendido antes.
        return vagas.order_by("criada_em" if situacao == Vaga.Status.PENDENTE else "-criada_em")


@api_view(["POST"])
@permission_classes([PodeAprovar])
def avaliar_vagas(request):
    """Aprova ou recusa uma ou várias vagas de uma vez. Fica registrado quem decidiu e quando."""
    dados = AvaliarSerializer(data=request.data)
    dados.is_valid(raise_exception=True)
    aprovar = dados.validated_data["decisao"] == "aprovar"
    alteradas = (
        Vaga.objects.filter(id__in=dados.validated_data["ids"])
        .exclude(status=Vaga.Status.ENCERRADA)
        .update(
            status=Vaga.Status.APROVADA if aprovar else Vaga.Status.RECUSADA,
            motivo_recusa="" if aprovar else dados.validated_data["motivo"],
            avaliada_por=request.user,
            avaliada_em=timezone.now(),
        )
    )
    return Response({"alteradas": alteradas})


@api_view(["GET"])
@permission_classes([PodeAprovar])
def numeros(request):
    vagas_por_area = Counter()
    for area, quantidade in Vaga.publicas().values_list("area", "quantidade"):
        vagas_por_area[area.strip().capitalize()] += quantidade
    interesses_por_area = Counter(
        area.strip().capitalize() for area in Interesse.objects.values_list("vaga__area", flat=True)
    )
    curriculos_por_area = Counter()
    for areas in Curriculo.objects.values_list("areas_interesse", flat=True):
        curriculos_por_area.update({a.strip().capitalize() for a in areas})

    todas = set(vagas_por_area) | set(interesses_por_area) | set(curriculos_por_area)
    por_area = sorted(
        (
            {
                "area": area,
                "vagas": vagas_por_area[area],
                "interesses": interesses_por_area[area],
                "curriculos": curriculos_por_area[area],
            }
            for area in todas
        ),
        key=lambda linha: (-(linha["vagas"] + linha["curriculos"]), linha["area"]),
    )
    publicas = Vaga.publicas()
    return Response(
        {
            "pendentes": Vaga.objects.filter(status=Vaga.Status.PENDENTE).count(),
            "vagas_abertas": publicas.count(),
            "postos_abertos": sum(publicas.values_list("quantidade", flat=True)),
            "curriculos": Curriculo.objects.count(),
            "empresas": Usuario.objects.filter(tipo=Usuario.Tipo.EMPRESA, is_active=True).count(),
            "interesses": Interesse.objects.count(),
            "por_area": por_area[:12],
        }
    )


class EmpresasView(generics.ListAPIView):
    permission_classes = [PodeAprovar]
    serializer_class = EmpresaPainelSerializer

    def get_queryset(self):
        empresas = Usuario.objects.filter(tipo=Usuario.Tipo.EMPRESA).annotate(
            vagas_abertas=Count("vagas", filter=Q(vagas__status=Vaga.Status.APROVADA)),
            vagas_total=Count("vagas"),
        )
        if busca := self.request.query_params.get("busca", "").strip():
            empresas = empresas.filter(Q(nome__icontains=busca) | Q(documento__contains=busca))
        return empresas.order_by("-criado_em")


@api_view(["POST"])
@permission_classes([PodeAprovar])
def bloquear_empresa(request, pk):
    """Empresa bloqueada não entra mais e as vagas dela somem da lista pública. Dá para desbloquear."""
    empresa = get_object_or_404(Usuario, pk=pk, tipo=Usuario.Tipo.EMPRESA)
    empresa.is_active = not bool(request.data.get("bloqueada", True))
    empresa.save(update_fields=["is_active"])
    return Response({"bloqueada": not empresa.is_active})


def _permissao():
    return Permission.objects.get(codename="aprovar_vaga", content_type__app_label="vagas")


class EquipeView(generics.ListCreateAPIView):
    """Quem pode aprovar. Qualquer aprovador pode incluir outra pessoa pelo CPF."""

    permission_classes = [PodeAprovar]
    pagination_class = None
    serializer_class = MembroSerializer

    def get_queryset(self):
        return Usuario.objects.filter(user_permissions=_permissao(), is_active=True).order_by("nome")

    def create(self, request, *args, **kwargs):
        dados = NovoMembroSerializer(data=request.data)
        dados.is_valid(raise_exception=True)
        usuario = dados.validated_data["documento"]
        usuario.tipo = Usuario.Tipo.PREFEITURA
        usuario.save(update_fields=["tipo"])
        usuario.user_permissions.add(_permissao())
        return Response(MembroSerializer(usuario).data, status=status.HTTP_201_CREATED)


@api_view(["DELETE"])
@permission_classes([PodeAprovar])
def remover_da_equipe(request, pk):
    if pk == request.user.pk:
        return Response({"detail": "Você não pode tirar a sua própria permissão."}, status=status.HTTP_400_BAD_REQUEST)
    usuario = get_object_or_404(Usuario, pk=pk, user_permissions=_permissao())
    usuario.user_permissions.remove(_permissao())
    usuario.tipo = Usuario.Tipo.CANDIDATO
    usuario.save(update_fields=["tipo"])
    return Response(status=status.HTTP_204_NO_CONTENT)


def _planilha(nome, cabecalho, linhas):
    livro = Workbook()
    folha = livro.active
    folha.title = nome.capitalize()
    folha.append(cabecalho)
    for celula in folha[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="1D4ED8")
    for linha in linhas:
        folha.append(linha)
    for coluna in folha.columns:
        largura = max(len(str(c.value or "")) for c in coluna)
        folha.column_dimensions[coluna[0].column_letter].width = min(max(largura + 2, 10), 50)
    folha.freeze_panes = "A2"
    saida = BytesIO()
    livro.save(saida)
    resposta = HttpResponse(
        saida.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    data = timezone.localdate().isoformat()
    resposta["Content-Disposition"] = f'attachment; filename="meuemprego-{nome}-{data}.xlsx"'
    return resposta


def _data(valor):
    return timezone.localtime(valor).strftime("%d/%m/%Y") if valor else ""


@api_view(["GET"])
@permission_classes([PodeAprovar])
def exportar(request, tipo):
    if tipo == "vagas":
        vagas = Vaga.objects.select_related("empresa").annotate(interessados=Count("interesses"))
        return _planilha(
            "vagas",
            ["Cargo", "Área", "Empresa", "CNPJ", "Bairro", "Salário", "Contrato", "Horário", "Vagas", "Situação",
             "Interessados", "Criada em"],
            (
                [v.cargo, v.area, v.empresa.nome, v.empresa.documento, v.bairro,
                 float(v.salario) if v.salario else "A combinar", v.get_contrato_display(), v.get_horario_display(),
                 v.quantidade, v.get_status_display(), v.interessados, _data(v.criada_em)]
                for v in vagas
            ),
        )  # fmt: skip
    if tipo == "curriculos":
        # Sem CPF: a planilha é para números e contato, e circula mais do que o sistema (LGPD).
        curriculos = Curriculo.objects.select_related("usuario").annotate(interesses=Count("usuario__interesses"))
        return _planilha(
            "curriculos",
            ["Nome", "WhatsApp", "Bairro", "Escolaridade", "Áreas", "CNH", "Interesses enviados", "Atualizado em"],
            (
                [c.usuario.nome, c.usuario.telefone, c.bairro, c.get_escolaridade_display(),
                 ", ".join(c.areas_interesse), c.categoria_cnh if c.tem_cnh else "Não", c.interesses,
                 _data(c.atualizado_em)]
                for c in curriculos
            ),
        )  # fmt: skip
    if tipo == "empresas":
        empresas = Usuario.objects.filter(tipo=Usuario.Tipo.EMPRESA).annotate(vagas_total=Count("vagas"))
        return _planilha(
            "empresas",
            ["Empresa", "CNPJ", "WhatsApp", "Vagas cadastradas", "Bloqueada", "Desde"],
            (
                [e.nome, e.documento, e.telefone, e.vagas_total, "Sim" if not e.is_active else "Não",
                 _data(e.criado_em)]
                for e in empresas
            ),
        )  # fmt: skip
    return Response({"detail": "Planilha não existe."}, status=status.HTTP_404_NOT_FOUND)
